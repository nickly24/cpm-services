"""Focused loopback template -> parse -> native preview acceptance.

Creates only synthetic fixtures using owner APIs, reuses an active fixture
direction, does no commit/grading and no direct database writes. Fixtures are
retained for review. Never emits credentials, bearer tokens or service keys.
"""
import importlib.util
from io import BytesIO
import json
import secrets
import urllib.error
import urllib.request

from openpyxl import load_workbook
import http_test as base

checks, requests, fixtures = [], 0, {"tests": [], "users": []}
run = secrets.token_hex(6)
report_path = base.ROOT / "dev/local/tests-import-templates-http-report.json"


def expect(ok, label):
    if not ok:
        raise AssertionError(label)
    checks.append(label)


def call(path, token=None, method="GET", body=None, code=200, origin=base.TESTS, headers=None):
    global requests
    headers = dict(headers or {})
    if token:
        headers["Authorization"] = "Bearer " + token
    data = None
    if body is not None:
        data = json.dumps(body, ensure_ascii=False).encode()
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(origin + path, data=data, method=method, headers=headers)
    requests += 1
    try:
        with urllib.request.urlopen(req, timeout=30) as res:
            status, raw, response_headers = res.status, res.read(), res.headers
    except urllib.error.HTTPError as res:
        status, raw, response_headers = res.code, res.read(), res.headers
    expect(status == code, f"{method} {path}: HTTP {status}, expected {code}")
    return raw, response_headers


def api(path, token, method="GET", body=None, code=200, origin=base.TESTS, headers=None):
    raw, h = call(path, token, method, body, code, origin, headers)
    result = json.loads(raw)
    return result.get("data", result), h


def download(path, token, mime, filename):
    raw, headers = call(path, token)
    expect(headers.get("Content-Type") == mime, "exact template MIME")
    expect(headers.get("Content-Disposition") == f'attachment; filename="{filename}"', "fixed attachment filename")
    expect(headers.get("Content-Length") == str(len(raw)), "actual file byte length")
    expect(headers.get("Cache-Control") == "no-store" and headers.get("X-Content-Type-Options") == "nosniff", "private download headers")
    return raw


def main():
    expect(base.TESTS == "http://127.0.0.1:8004" and base.CORE == "http://127.0.0.1:8003", "loopback only")
    for origin in [base.TESTS, base.AUTH, base.CORE]:
        call("/health/ready", origin=origin)
    admin, student = base.login("dev_admin"), base.login("dev_student_01")
    foreign = base.login("dev_admin_math", 2001, 2001)
    directions, _ = api("/v1/core/directions", admin, origin=base.CORE)
    direction = next(item for item in directions if item["id"] == 1056)
    private = {"X-Service-Key": base.read_env(base.ROOT / "dev/.env")["AUTH_INTERNAL_KEY"]}
    root = "/v1/tests/divisions/1001/imports/json/"
    for format, mime, filename in [
        ("cpm_json", "application/json; charset=utf-8", "primer_importa_testa_cpm.json"),
        ("online_test_pad", "text/javascript; charset=utf-8", "primer_importa_testa_online_test_pad.js"),
    ]:
        raw = download(root + "template?direction_id=1056&format=" + format, admin, mime, filename)
        if format == "cpm_json":
            payload = json.loads(raw)
            expect(payload["direction"] == direction["name"] and payload["published"] is False, "native sample current direction and hidden")
        else:
            text = raw.decode()
            source = json.loads(text.removeprefix("var test = ").strip().removesuffix(";"))
            expect(len(source["Questions"]) == 3, "OTP pure JSON assignment parse without executing JS")
            payload = {"importFormat": format, "sourceText": text, "metadata": {
                "direction": direction["name"], "startDate": "2099-01-01T10:00", "endDate": "2099-01-01T11:00", "published": False}}
        preview, _ = api("/internal/tests/imports/json/validate", admin, "POST", {
            "organization_id": 1001, "division_id": 1001, "direction_id": 1056, "payload": payload}, headers=private)
        expect(preview["valid"] and preview["summary"]["questions_total"] == 3, "downloaded sample passes actual native preview")
        schema, _ = api(root + "schema?direction_id=1056&format=" + format, admin)
        expect(schema["capability"] == "supported" and schema["direction_id"] == 1056, "scoped schema supported")
    for path, token, code in [
        (root+"template?direction_id=1056", student, 403), (root+"schema?direction_id=1056", student, 403),
        (root+"template?direction_id=1057", admin, 404), (root+"template?direction_id=1056", foreign, 404),
        (root+"template?direction_id=1056&format=html", admin, 400),
        (root+"template?direction_id=1056&direction_id=1056", admin, 400),
        (root+"template?direction_id=1056", None, 401),
    ]:
        _, h = call(path, token, code=code)
        expect(h.get("Content-Disposition") is None, "error never downloaded as attachment")

    target, _ = api("/v1/tests/divisions/1001/external-tests", admin, "POST", {
        "name": "DEV Template " + run, "direction_id": 1056, "date": "2099-01-01"}, 201, headers=base.key())
    fixtures["tests"].append({"id": target["id"], "organization_id": 1001, "division_id": 1001})
    prefix = f'/v1/tests/divisions/1001/external-tests/{target["id"]}/imports/results/'
    raw = download(prefix+"template", admin,
                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "external-test-results-template.xlsx")
    wb = load_workbook(BytesIO(raw))
    sheet = wb.active
    expect(sheet.max_column == 8 and sheet["G1"].value == "account_id" and sheet["H1"].value == "legacy_student_id", "all legacy and optional ID columns")
    expect(sheet["B2"].value == 0 and sheet["B3"].value == 87.5, "XLSX real numeric zero and decimal")
    parser_path = base.ROOT / "cpm-back-main/cpm_back/services/external_test_results_import/parse_excel.py"
    spec = importlib.util.spec_from_file_location("legacy_template_parser", parser_path)
    parser = importlib.util.module_from_spec(spec); spec.loader.exec_module(parser)
    parsed = parser.parse_external_results_excel(raw)
    expect(parsed["status"] and len(parsed["rows"]) == 2, "old XLSX parser understands downloaded workbook")
    sample_raw = download(prefix+"sample", admin, "application/json; charset=utf-8", "external-test-results-sample.json")
    sample = json.loads(sample_raw)
    expect(sample["test_id"] == target["id"] and sample["expected_test_revision"] == target["revision"], "sample uses selected scoped target/revision")
    preview, _ = api("/internal/tests/external-results/validate", admin, "POST", sample, headers=private)
    expect(not preview["can_commit"], "fictional sample does not impersonate actual students")
    for index in range(2):
        name = "Шаблонов" + run + " " + ("Первый" if index == 0 else "Второй")
        created, _ = api("/v1/core/divisions/1001/users", admin, "POST", {
            "role": "student", "full_name": name, "class": 10, "login": "dev_tpl_" + run + "_" + str(index), "password": base.password}, 201, origin=base.CORE)
        profile, _ = api("/v1/core/divisions/1001/users/" + str(created["account_id"]), admin, origin=base.CORE)
        person = {"account_id": created["account_id"], "student_id": profile["profile_id"], "full_name": name}
        fixtures["users"].append(person)
        rownum = index + 2
        sheet.cell(rownum, 1).value = name
        sheet.cell(rownum, 7).value = str(person["account_id"])
        sheet.cell(rownum, 8).value = str(person["student_id"])
    filled = BytesIO(); wb.save(filled); wb.close()
    parsed = parser.parse_external_results_excel(filled.getvalue())
    sample["rows"] = [{"row_id": "sheet-1-row-"+str(row["row"]), "row": row["row"], "full_name": row["full_name"], "percent": row["percent"]} for row in parsed["rows"]]
    for row, person in zip(sample["rows"], fixtures["users"]):
        row.update(account_id=person["account_id"], legacy_student_id=person["student_id"])
    preview, _ = api("/internal/tests/external-results/validate", admin, "POST", sample, headers=private)
    expect(preview["can_commit"] and preview["summary"]["import_rows"] == 2, "filled downloaded workbook rows validated in actual Core scope")
    expect([row["percent"] for row in preview["rows"]] == [0, 87.5], "native preview preserves zero and decimal")
    mismatch = json.loads(json.dumps(sample)); mismatch["rows"][0]["legacy_student_id"] = fixtures["users"][1]["student_id"]
    bad, _ = api("/internal/tests/external-results/validate", admin, "POST", mismatch, headers=private)
    expect(not bad["can_commit"] and any(err["code"] == "student_account_mismatch" for err in bad["rows"][0]["errors"]), "different account/profile rejects without name fallback")
    for suffix in ["template", "sample", "schema"]:
        call(prefix+suffix, student, code=403)
        call(prefix+suffix, foreign, code=404)
        call(prefix+suffix+"?target_id="+target["id"], admin, code=400)
    schema, _ = api(prefix+"schema", admin)
    percent = schema["validate_schema"]["properties"]["rows"]["items"]["properties"]["percent"]
    expect(percent["type"] == "number" and percent["minimum"] == 0 and percent["maximum"] == 100, "schema forbids null percent")
    after, _ = api(f'/v1/tests/divisions/1001/external-tests/{target["id"]}', admin)
    expect(after["revision"] == target["revision"], "download/parse/preview never changed target revision")


if __name__ == "__main__":
    status, error = "passed", None
    try:
        main()
    except Exception as exc:
        status, error = "failed", str(exc)
    finally:
        closed = 0
        for token in base.sessions:
            try:
                call("/v1/auth/logout", token, "POST", origin=base.AUTH)
                closed += 1
            except Exception:
                status = "failed"; error = error or "logout failed"
        report = {"status": status, "run": run, "http_requests": requests + len(base.sessions),
                  "tracked_http_requests": requests, "login_http_requests": len(base.sessions), "semantic_checks": len(checks),
                  "checks": checks, "fixtures": fixtures, "fixtures_retained": True,
                  "closed_sessions": closed, "directions_mutated": False,
                  "grading_or_result_commit": False, "direct_database_writes": False, "error": error}
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
        print(json.dumps({key: report[key] for key in ["status", "http_requests", "semantic_checks", "closed_sessions", "error"]}, ensure_ascii=False))
        if status != "passed":
            raise SystemExit(1)
